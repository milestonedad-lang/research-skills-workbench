"""Bounded research agent with an auditable JSON tool protocol."""
import asyncio
import json
import math
import os
import shutil
import time
from urllib.parse import quote

import httpx

from server.core import BudgetExceeded, digest, now, safe_path, uid
from server.providers import completion, load_provider, prices

PROTOCOL = """你是科研助手。依据真实输入与工具返回工作，区分事实、推断和未完成事项，不编造文献、数据或实验结果。
每次只输出一个 JSON 对象，不使用 Markdown 代码围栏：
{"action":"read_file","path":"文件名"} 读取项目原始资料；
{"action":"read_resource","path":"references/文件名"} 读取所选技能资源；
{"action":"crossref_search","query":"检索式"} 检索真实文献；
{"action":"crossref_doi","doi":"10.xxx/xxx"} 核验 DOI 元数据；
{"action":"python","code":"Python 代码"} 在 Docker 运行；输入在 /inputs，技能在 /skills，输出必须写到 /outputs。
{"action":"write_file","path":"filename.md","content":"文字"} 保存文本；
{"action":"final","content":"完整中文报告"} 完成。数据分析与绘图须实际执行代码，不能只声称已执行。
Python 环境包含 pandas,numpy,scipy,scikit-learn,statsmodels,matplotlib,seaborn,openpyxl,python-docx,pypdf,reportlab。
代码无法访问模型密钥，默认不能联网。联网文献查询使用上述工具。最终报告说明方法、证据及局限。
"""


def prepare_run(store, payload, budget_id=None, frozen_provider=None, frozen_settings=None):
    store.get("projects", payload["project_id"])
    try:
        stored_provider = frozen_provider or store.get('providers', payload.get('provider_id'))
    except KeyError:
        raise ValueError('尚未配置所选模型，请先到设置中添加') from None
    provider = dict(stored_provider)
    if provider.get('api_key'):
        from server.providers import cipher
        provider['api_key'] = cipher(store).decrypt(provider['api_key'].encode()).decode()
    if provider['kind'] != 'ollama' and not provider.get('api_key'):
        raise ValueError('此模型需要 API 密钥')
    prices(provider)
    prompt = payload.get("prompt", "").strip()
    if not prompt or len(prompt) > 50000:
        raise ValueError("任务描述为空或超过限制")
    budget = float(payload.get("budget", 0))
    if not math.isfinite(budget) or budget < 0 or (provider["kind"] != "ollama" and budget <= 0):
        raise ValueError("付费运行必须设置正数预算上限")
    version_id = payload.get("version_id")
    skill_id = payload.get("skill_id")
    if not version_id and skill_id:
        scope = payload["project_id"]
        installations = store.installations()
        selected = next((i for i in installations if i["skill_id"] == skill_id and i["scope"] == scope), None)
        selected = selected or next((i for i in installations if i["skill_id"] == skill_id and i["scope"] == "global"), None)
        if not selected:
            raise ValueError("请先安装技能固定版本")
        version_id = selected["version_id"]
    skill = store.version(version_id) if version_id else None
    if skill:
        store.install(skill["id"], version_id, scope="cache")
        skill_id = skill["id"]
    identity = uid()
    output = store.root / "projects" / payload["project_id"] / "runs" / identity
    output.mkdir(parents=True)
    inputs = store.root / "projects" / payload["project_id"] / "inputs"
    # Snapshot inputs: later uploads cannot change an active comparison.
    frozen = output / "inputs"
    shutil.copytree(inputs, frozen)
    manifests = {str(p.relative_to(frozen)).replace("\\", "/"): __import__('hashlib').sha256(p.read_bytes()).hexdigest() for p in frozen.rglob("*") if p.is_file()}
    try:
        config = store.get("settings", "main")
    except KeyError:
        config = {}
    if frozen_settings is not None:
        config = frozen_settings
    if not budget_id:
        budget_id = identity
        store.create_budget(budget_id, budget)
    # Provider configuration is immutable for this run, encrypted key included internally.
    store.save("run-providers", {**stored_provider, "id": identity})
    r = {"id": identity, "project_id": payload["project_id"], "prompt": prompt, "provider_id": provider["id"],
         "model": provider["model"], "actual_model": None, "skill_id": skill_id, "version_id": version_id,
         "input_hashes": manifests, "input_digest": digest(json.dumps(manifests, sort_keys=True)),
         "status": "queued", "created_at": now(), "budget_id": budget_id,
         "max_steps": min(max(int(payload.get("max_steps", 12)), 1), 30),
         "max_output": min(max(int(payload.get("max_output", 2048)), 128), 8192),
         "temperature": min(max(float(payload.get("temperature", 0)), 0), 2),
         "runner_image": config.get("runner_image", "research-workbench-runner:0.1.0"),
         "runner_digest": config.get("runner_digest"),
         "runner_policy": "network-none,inputs-readonly,cpu-2,memory-2g", "events": [], "tokens": {"input": 0, "output": 0},
         "cost": 0, "duration": 0, "artifacts": [], "benchmark_id": payload.get("benchmark_id")}
    return store.save("runs", r)


def public_provider_for_run(store, identity):
    from server.providers import cipher
    provider = store.get("run-providers", identity)
    if provider.get("api_key"):
        provider["api_key"] = cipher(store).decrypt(provider["api_key"].encode()).decode()
    return provider


async def billed_completion(store, provider, budget_id, system, messages, max_output, temperature):
    input_price, output_price = prices(provider)
    # UTF-8 byte length is a conservative bound for byte-based tokenizers; allowance for protocol overhead.
    input_bound = len(json.dumps({"system": system, "messages": messages}, ensure_ascii=False).encode()) + 4096
    if input_bound > 200000:
        raise ValueError("任务上下文超过上限，请缩小资料范围")
    reservation = (input_bound * input_price + max_output * output_price) / 1_000_000
    store.reserve(budget_id, reservation)
    try:
        text, input_tokens, output_tokens, actual_model = await completion(provider, system, messages, max_output, temperature)
    except BaseException:
        # A timeout/cancellation may already have been billed; charge the reservation conservatively.
        store.settle(budget_id, reservation, reservation)
        raise
    known = input_tokens is not None and output_tokens is not None
    actual = (input_tokens * input_price + output_tokens * output_price) / 1_000_000 if known else reservation
    store.settle(budget_id, reservation, actual)
    return text, {"input": input_tokens or 0, "output": output_tokens or 0}, actual, actual_model, known


async def run_python(store, run, code):
    if not shutil.which("docker"):
        raise ValueError("未检测到 Docker：安装并启动 Docker Desktop 后，构建 runner 镜像再运行代码任务")
    directory = store.root / "projects" / run["project_id"] / "runs" / run["id"]
    output = directory / "outputs"
    output.mkdir(exist_ok=True)
    script = output / ("analysis-" + uid()[:8] + ".py")
    script.write_text(code, encoding="utf-8")
    user = '65534:65534'
    if hasattr(os, 'getuid'):
        if os.getuid() == 0:
            os.chown(output, 65534, 65534)
        else:
            user = f'{os.getuid()}:{os.getgid()}'
    name = "research-" + run["id"][:20]
    # Resolve the tag to a fixed image ID once; subsequent steps use the same image.
    if not run.get("runner_digest"):
        inspect = await asyncio.create_subprocess_exec("docker", "image", "inspect", run["runner_image"], "--format", "{{.Id}}", stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
        out, _ = await asyncio.wait_for(inspect.communicate(), 20)
        if inspect.returncode:
            raise ValueError("执行镜像不存在：请运行 docker build -t research-workbench-runner:0.1.0 -f docker/runner.Dockerfile .")
        run["runner_digest"] = out.decode().strip()
    args = ["docker", "run", "--rm", "--name", name, "--network", "none", "--cpus", "2", "--memory", "2g", "--pids-limit", "128",
            "--read-only", "--cap-drop", "ALL", "--security-opt", "no-new-privileges", "--user", user,
            "--tmpfs", "/tmp:rw,noexec,nosuid,size=128m", "--mount", f"type=bind,src={directory / 'inputs'},dst=/inputs,readonly",
            "--mount", f"type=bind,src={output},dst=/outputs", "--workdir", "/outputs", "-e", "MPLCONFIGDIR=/tmp/matplotlib"]
    if run.get("version_id"):
        args.extend(["--mount", f"type=bind,src={store.root / 'installed' / run['version_id']},dst=/skills,readonly"])
    args.extend([run["runner_digest"], "python", "/outputs/" + script.name])
    process = await asyncio.create_subprocess_exec(*args, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
    async def drain(stream, limit):
        kept = bytearray()
        while chunk := await stream.read(32768):
            kept.extend(chunk[:max(0, limit-len(kept))])
        return bytes(kept)
    async def collect():
        stdout, stderr, _ = await asyncio.gather(drain(process.stdout, 50000), drain(process.stderr, 15000), process.wait())
        return stdout, stderr
    try:
        stdout, stderr = await asyncio.wait_for(collect(), 120)
    except BaseException:
        cleanup = await asyncio.create_subprocess_exec("docker", "rm", "-f", name, stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
        await cleanup.wait()
        if process.returncode is None:
            process.kill()
            await process.wait()
        raise
    return {"exit_code": process.returncode, "stdout": stdout.decode(errors="replace")[:50000], "stderr": stderr.decode(errors="replace")[:15000], "script": script.name}


async def tool_call(store, run, action):
    directory = store.root / "projects" / run["project_id"] / "runs" / run["id"]
    kind = action.get("action")
    if kind in {"read_file", "read_resource"}:
        root = directory / "inputs" if kind == "read_file" else store.root / "installed" / (run.get("version_id") or "none")
        path = safe_path(root, action["path"])
        if path.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            return {"text": "\n".join(p.extract_text() or "" for p in reader.pages[:50])[:40000], "pages": len(reader.pages), "truncated": len(reader.pages) > 50}
        if path.suffix.lower() == ".docx":
            from docx import Document
            return {"text": "\n".join(p.text for p in Document(path).paragraphs)[:40000]}
        return {"text": path.read_text(encoding="utf-8", errors="replace")[:40000]}
    if kind == "write_file":
        output = directory / "outputs"
        output.mkdir(exist_ok=True)
        path = safe_path(output, action["path"])
        content = str(action.get("content", ""))
        if len(content) > 500000:
            raise ValueError("输出文件过大")
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("x", encoding="utf-8") as f:
            f.write(content)
        return {"written": path.name}
    if kind == "python":
        return await run_python(store, run, action["code"])
    if kind in {"crossref_search", "crossref_doi"}:
        url = "https://api.crossref.org/works"
        params = {"query": str(action.get("query", ""))[:500], "rows": 8} if kind == "crossref_search" else None
        if kind == "crossref_doi":
            url += "/" + quote(str(action["doi"]).removeprefix("https://doi.org/"), safe="")
        async with httpx.AsyncClient(timeout=25) as client:
            response = await client.get(url, params=params, headers={"User-Agent": "research-workbench/0.1"})
            if response.is_error:
                raise ValueError(f"Crossref 返回 HTTP {response.status_code}")
            message = response.json()["message"]
        items = message.get("items", []) if kind == "crossref_search" else [message]
        return {"source": url, "retrieved_at": now(), "records": [{k: i.get(k) for k in ("DOI", "title", "author", "published", "URL", "type")} for i in items]}
    raise ValueError("不支持的工具动作")


def save_reports(directory, content):
    from docx import Document
    output = directory / "outputs"
    output.mkdir(exist_ok=True)
    (output / "report.md").write_text(content, encoding="utf-8")
    doc = Document()
    for line in content.splitlines():
        if line.startswith("#"):
            doc.add_heading(line.lstrip("# "), level=min(len(line) - len(line.lstrip("#")), 3))
        else:
            doc.add_paragraph(line)
    doc.save(output / "report.docx")


async def execute_run(store, identity):
    run = store.get("runs", identity)
    started = time.monotonic()
    directory = store.root / "projects" / run["project_id"] / "runs" / identity
    run.update(status="running", started_at=now())
    run["events"].append({"type": "started", "at": now()})
    store.save("runs", run)
    try:
        provider = public_provider_for_run(store, identity)
        skill = store.version(run["version_id"]) if run.get("version_id") else None
        system = PROTOCOL + ("\n所选技能：\n" + skill["body"][:40000] if skill else "\n无技能基线。")
        files = [{"name": str(p.relative_to(directory / "inputs")), "size": p.stat().st_size} for p in (directory / "inputs").rglob("*") if p.is_file()]
        messages = [{"role": "user", "content": run["prompt"] + "\n原始文件：" + json.dumps(files, ensure_ascii=False)}]
        for _ in range(run["max_steps"]):
            if store.get("runs", identity).get("cancel_requested"):
                raise asyncio.CancelledError()
            text, tokens, cost, actual_model, known = await billed_completion(store, provider, run["budget_id"], system, messages, run["max_output"], run["temperature"])
            run["cost"] += cost
            run["actual_model"] = actual_model
            run["usage_known"] = known
            for key in tokens:
                run["tokens"][key] += tokens[key]
            run["events"].append({"type": "model", "content": text, "tokens": tokens, "at": now()})
            store.save("runs", run)
            try:
                clean = text.strip()
                if clean.startswith("```"):
                    clean = clean.split("\n", 1)[-1].rsplit("```", 1)[0]
                action = json.loads(clean)
                if not isinstance(action, dict):
                    raise ValueError()
            except (ValueError, TypeError):
                action = {"action": "final", "content": text}
            if action.get("action") == "final":
                content = str(action.get("content", ""))
                if not content.strip():
                    raise ValueError("模型返回空结果")
                save_reports(directory, content)
                run.update(status="completed", output=content)
                break
            messages.append({"role": "assistant", "content": text})
            try:
                result = await tool_call(store, run, action)
            except (ValueError, OSError, httpx.HTTPError) as exc:
                result = {"error": str(exc)[:400]}
                run["tool_failures"] = run.get("tool_failures", 0) + 1
            run["events"].append({"type": "tool", "action": action.get("action"), "request": action, "result": result, "at": now()})
            messages.append({"role": "user", "content": "工具实际返回：" + json.dumps(result, ensure_ascii=False)})
            store.save("runs", run)
        else:
            raise ValueError("达到最大执行步数，任务尚未完成")
    except asyncio.CancelledError:
        run.update(status="cancelled", error="运行已取消；已发送请求的费用可能仍由服务商计费")
    except BudgetExceeded as exc:
        run.update(status='budget_exhausted', error=str(exc))
    except Exception as exc:
        run.update(status="failed", error=str(exc)[:400])
    finally:
        run.update(finished_at=now(), duration=time.monotonic() - started)
        run["budget"] = store.budget(run["budget_id"])
        output = directory / "outputs"
        run["artifacts"] = [{"name": str(p.relative_to(directory.parent.parent)).replace("\\", "/"), "sha256": __import__('hashlib').sha256(p.read_bytes()).hexdigest(), "size": p.stat().st_size}
                            for p in output.rglob("*") if p.is_file()] if output.exists() else []
        run["events"].append({"type": "finished", "status": run["status"], "at": now()})
        store.save("runs", run)
        (directory / "trace.json").write_text(json.dumps(run, ensure_ascii=False, indent=2), encoding="utf-8")
    return run


async def translate_skills(store, payload):
    provider = load_provider(store, payload["provider_id"])
    prices(provider)
    identity = uid()
    store.create_budget(identity, float(payload.get("budget", 0)))
    results = []
    for skill_id in payload.get("skill_ids", [])[:30]:
        skill = store.skill(skill_id)
        text, _, _, _, _ = await billed_completion(store, provider, identity, "将科研技能用途译为简洁准确的中文，保留适用边界，仅输出译文。", [{"role":"user", "content": skill["description"]}], 512, 0)
        with store.connect() as c:
            c.execute("UPDATE skills SET description_zh=? WHERE id=?", (text.strip(), skill_id))
            record = store.version(skill["version_id"])
            record.update(description_zh=text.strip(), translation_status="AI 翻译，待复核")
            c.execute("UPDATE versions SET record=? WHERE id=?", (json.dumps(record, ensure_ascii=False), skill["version_id"]))
        results.append({"skill_id":skill_id, "translation":text})
    return {"items":results, "budget":store.budget(identity)}
